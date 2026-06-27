from django.contrib import admin
from django.urls import path
from django.shortcuts import render, redirect
from django.contrib import messages
from django import forms
import openpyxl
from orders.models import CargoRequest
from .models import Province, City, Country, DestinationCity, Port

# ==========================================
# Forms
# ==========================================

class ExcelImportForm(forms.Form):
    excel_file = forms.FileField(label="فایل اکسل (ستون A: استان، ستون B: شهر)")

class CountryExcelImportForm(forms.Form):
    excel_file = forms.FileField(label="فایل اکسل (کشور، شهر، پورت)")


# ==========================================
# Inlines
# ==========================================

class CityInline(admin.TabularInline):
    model = City
    extra = 0

class DestCityInline(admin.TabularInline):
    model = DestinationCity
    extra = 0

class PortInline(admin.TabularInline):
    model = Port
    extra = 0


# ==========================================
# Province & City Admin
# ==========================================

@admin.register(Province)
class ProvinceAdmin(admin.ModelAdmin):
    list_display = ["name", "is_active"]
    inlines = [CityInline]
    search_fields = ["name"]

    
    # معرفی قالب اختصاصی برای اضافه کردن دکمه آپلود
    change_list_template = "admin/locations/province/change_list.html"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.admin_site.admin_view(self.import_excel), name='locations_province_import_excel'),
        ]
        # حتما custom_urls اول باشد
        return custom_urls + urls

    def import_excel(self, request):
        if request.method == "POST":
            form = ExcelImportForm(request.POST, request.FILES)
            if form.is_valid():
                excel_file = request.FILES["excel_file"]
                try:
                    wb = openpyxl.load_workbook(excel_file, data_only=True)
                    sheet = wb.active
                    
                    # فرض می‌کنیم سطر اول هدر است، پس از سطر دوم شروع می‌کنیم
                    count = 0
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        province_name = row[0]
                        city_name = row[1]
                        
                        if province_name and city_name:
                            prov_name = str(province_name).strip()
                            cit_name = str(city_name).strip()
                            
                            # ایجاد یا دریافت استان
                            province_obj, _ = Province.objects.get_or_create(name=prov_name)
                            # ایجاد یا دریافت شهر
                            City.objects.get_or_create(province=province_obj, name=cit_name)
                            
                            count += 1
                            
                    messages.success(request, f"اطلاعات با موفقیت پردازش شد. {count} ردیف بررسی/ثبت گردید.")
                    return redirect("..")
                except Exception as e:
                    messages.error(request, f"خطا در پردازش فایل: {e}")
        else:
            form = ExcelImportForm()
            
        context = dict(
            self.admin_site.each_context(request),
            form=form,
            opts=self.model._meta,
            title="آپلود فایل اکسل استان و شهر",
        )
        return render(request, "admin/locations/excel_form.html", context)

# ثبت مدل City به صورت ساده
@admin.register(City)
class CityAdmin(admin.ModelAdmin):
    list_display = ["name", "province", "is_active"]
    list_filter = ["province", "is_active"]

    search_fields = [
        "name",
        "province__name",
    ]

    autocomplete_fields = [
        "province",
    ]


# ==========================================
# Country, DestinationCity & Port Admin
# ==========================================

# نگاشت (Mapping) برای تشخیص نوع پورت از روی متن فارسی یا انگلیسی
PORT_TYPE_MAP = {
    "بندر": Port.PortType.SEA,
    "sea": Port.PortType.SEA,
    "فرودگاه": Port.PortType.AIR,
    "air": Port.PortType.AIR,
    "گمرک زمینی": Port.PortType.land,
    "land": Port.PortType.land,
    "گمرک ریلی": Port.PortType.rail,
    "rail": Port.PortType.rail,
}

@admin.register(Country)
class CountryAdmin(admin.ModelAdmin):
    list_display = ["name", "code", "is_active"]
    inlines = [DestCityInline]
    search_fields = ['name', 'iso_code']
    
    # معرفی قالب اختصاصی برای اضافه کردن دکمه آپلود
    change_list_template = "admin/locations/country/change_list.html"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.admin_site.admin_view(self.import_excel), name='locations_country_import_excel'),
        ]
        # حتما custom_urls اول باشد
        return custom_urls + urls

    def import_excel(self, request):
        if request.method == "POST":
            form = CountryExcelImportForm(request.POST, request.FILES)
            if form.is_valid():
                excel_file = request.FILES["excel_file"]
                try:
                    wb = openpyxl.load_workbook(excel_file, data_only=True)
                    sheet = wb.active
                    
                    count = 0
                    # فرض: سطر اول هدر است.
                    # ستون‌ها: 0:Country, 1:CountryCode, 2:City, 3:PortType, 4:PortName, 5:PortCode
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        country_name = str(row[0]).strip() if row[0] else None
                        country_code = str(row[1]).strip()[:3] if row[1] else None
                        city_name = str(row[2]).strip() if row[2] else None
                        port_type_raw = str(row[3]).strip() if row[3] else None
                        port_name = str(row[4]).strip() if row[4] else None
                        port_code = str(row[5]).strip() if row[5] else ""
                        
                        if country_name and country_code and city_name and port_name and port_type_raw:
                            # 1. ایجاد یا دریافت کشور
                            country_obj, _ = Country.objects.get_or_create(
                                name=country_name,
                                defaults={'code': country_code}
                            )
                            
                            # 2. ایجاد یا دریافت شهر مقصد
                            city_obj, _ = DestinationCity.objects.get_or_create(
                                country=country_obj, 
                                name=city_name
                            )
                            
                            # 3. استخراج نوع پورت
                            p_type = PORT_TYPE_MAP.get(port_type_raw.lower(), Port.PortType.SEA)
                            
                            # 4. ایجاد یا دریافت پورت
                            Port.objects.get_or_create(
                                city=city_obj,
                                name=port_name,
                                defaults={
                                    'port_type': p_type,
                                    'code': port_code
                                }
                            )
                            
                            count += 1
                            
                    messages.success(request, f"عملیات موفق. {count} ردیف پورت/شهر/کشور ثبت یا بررسی شد.")
                    return redirect("..")
                except Exception as e:
                    messages.error(request, f"خطا در پردازش فایل: {e}")
        else:
            form = CountryExcelImportForm()
            
        context = dict(
            self.admin_site.each_context(request),
            form=form,
            opts=self.model._meta,
            title="آپلود اکسل مقاصد بین‌المللی",
        )
        return render(request, "admin/locations/country_excel_form.html", context)

@admin.register(DestinationCity)
class DestinationCityAdmin(admin.ModelAdmin):
    list_display = ["name", "country", "is_active"]
    inlines = [PortInline]

    search_fields = [
        "name",
        "country__name",
    ]

    autocomplete_fields = [
        "country",
    ]

@admin.register(Port)
class PortAdmin(admin.ModelAdmin):
    list_display = ["name", "port_type", "city", "code", "is_active"]
    list_filter = ["port_type", "is_active"]

    search_fields = [
        "name",
        "code",
        "city__name",
        "city__country__name",
    ]

    autocomplete_fields = [
        "city",
    ]
